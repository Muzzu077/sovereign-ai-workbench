"""
DOCX document processor.

Uses ``python-docx`` to extract paragraph text from ``.docx`` files
and converts them into the normalized ``Document`` representation.
"""

from __future__ import annotations

import logging
from pathlib import Path

import docx

from app.documents.models import (
    Document,
    DocumentMetadata,
    DocumentPage,
    ExtractionStatus,
    FileType,
)
from app.documents.processor import DocumentProcessor, ProcessingError

logger = logging.getLogger(__name__)


class DocxProcessor(DocumentProcessor):
    """Processor for DOCX files using python-docx."""

    @property
    def supported_extensions(self) -> list[str]:
        return ["docx"]

    def process(
        self,
        file_path: Path,
        document_id: str,
        *,
        max_pages: int = 200,
        max_characters: int = 500_000,
    ) -> Document:
        if not file_path.exists():
            raise ProcessingError(f"File not found: {file_path.name}")

        file_size = file_path.stat().st_size

        try:
            doc = docx.Document(str(file_path))
        except Exception as exc:
            raise ProcessingError(
                f"Cannot read DOCX '{file_path.name}': {exc}"
            ) from exc

        paragraphs: list[str] = []
        total_chars = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            remaining = max_characters - total_chars
            if remaining <= 0:
                break
            if len(text) > remaining:
                text = text[:remaining]

            paragraphs.append(text)
            total_chars += len(text)

        full_text = "\n".join(paragraphs)

        # DOCX doesn't have strict page boundaries accessible via python-docx
        # without complex section/break parsing. We treat the whole document
        # as a single page for now and note the paragraph count in metadata.
        page = DocumentPage(
            page_number=1,
            text=full_text,
            source="text_extraction",
        )

        # Extract available metadata from the core properties
        core = doc.core_properties
        author = core.author if core.author else None
        title = core.title if core.title else None

        return Document(
            document_id=document_id,
            filename=file_path.name,
            file_type=FileType.DOCX,
            file_size=file_size,
            page_count=1,
            extraction_status=ExtractionStatus.TEXT_EXTRACTED,
            text=full_text,
            pages=[page],
            metadata=DocumentMetadata(
                original_filename=file_path.name,
                stored_filename=file_path.name,
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                author=author,
                title=title,
                extra={"paragraph_count": len(paragraphs)},
            ),
        )

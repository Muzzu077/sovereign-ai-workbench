"""
Export service — converts markdown to styled DOCX documents.

Handles the subset of markdown produced by the local LLM:
  # / ## / ###   → Word headings (levels 1-3)
  **text**       → bold runs
  *text*         → italic runs
  `code`         → monospaced inline runs
  - / * bullet   → unordered list items
  1. ordered     → ordered list items
  ---            → horizontal rule (border)
  blank lines    → paragraph breaks

All processing runs entirely on-premise using python-docx.
No external network calls.
"""

from __future__ import annotations

import logging
import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)

# ── Regex patterns ───────────────────────────────────────────────────────────

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)")
_ORDERED_RE = re.compile(r"^\d+\.\s+(.*)")
_HR_RE = re.compile(r"^-{3,}$")

_INLINE_RE = re.compile(
    r"(\*\*(?P<bold>.+?)\*\*"
    r"|\*(?P<italic>.+?)\*"
    r"|`(?P<code>.+?)`"
    r"|\[(?P<link_text>.+?)\]\(.+?\))"
)


# ── Internal helpers ─────────────────────────────────────────────────────────


def _add_horizontal_rule(doc: Document) -> None:
    """Insert a thin bottom-border paragraph that mimics an <hr>."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _add_inline_styled_runs(para, raw_text: str) -> None:
    """Parse inline markdown into styled runs on an existing paragraph."""
    pos = 0
    for m in _INLINE_RE.finditer(raw_text):
        start, end = m.span()
        if start > pos:
            run = para.add_run(raw_text[pos:start])
            run.font.size = Pt(11)
        if m.group("bold"):
            run = para.add_run(m.group("bold"))
            run.bold = True
            run.font.size = Pt(11)
        elif m.group("italic"):
            run = para.add_run(m.group("italic"))
            run.italic = True
            run.font.size = Pt(11)
        elif m.group("code"):
            run = para.add_run(m.group("code"))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)
        elif m.group("link_text"):
            run = para.add_run(m.group("link_text"))
            run.underline = True
            run.font.size = Pt(11)
        pos = end
    if pos < len(raw_text):
        run = para.add_run(raw_text[pos:])
        run.font.size = Pt(11)


def _add_styled_paragraph(doc: Document, raw_text: str) -> None:
    """Add a normal paragraph with inline markdown styling."""
    p = doc.add_paragraph()
    _add_inline_styled_runs(p, raw_text)


# ── Public service class ─────────────────────────────────────────────────────


class ExportService:
    """Converts markdown text to binary DOCX content."""

    # Brand accent colour (teal from our design system)
    _ACCENT = RGBColor(0x0F, 0x76, 0x6E)

    def markdown_to_docx(self, title: str, markdown_text: str) -> bytes:
        """Convert markdown to a styled DOCX document.

        Args:
            title:         Document title (Heading 0).
            markdown_text: Raw markdown string.

        Returns:
            Raw bytes of the .docx file.
        """
        doc = Document()

        # Document-level styles
        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)

        # Document title
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in title_para.runs:
            run.font.color.rgb = self._ACCENT

        doc.add_paragraph()  # spacing after title

        lines = markdown_text.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i]

            # Blank line
            if not line.strip():
                i += 1
                continue

            # Horizontal rule
            if _HR_RE.match(line):
                _add_horizontal_rule(doc)
                i += 1
                continue

            # Heading
            h_match = _HEADING_RE.match(line)
            if h_match:
                level = len(h_match.group(1))
                heading_text = h_match.group(2).strip()
                h = doc.add_heading(heading_text, level=level)
                for run in h.runs:
                    run.font.color.rgb = self._ACCENT
                i += 1
                continue

            # Bullet list item
            b_match = _BULLET_RE.match(line)
            if b_match:
                p = doc.add_paragraph(style="List Bullet")
                _add_inline_styled_runs(p, b_match.group(1))
                i += 1
                continue

            # Ordered list item
            o_match = _ORDERED_RE.match(line)
            if o_match:
                p = doc.add_paragraph(style="List Number")
                _add_inline_styled_runs(p, o_match.group(1))
                i += 1
                continue

            # Normal paragraph
            _add_styled_paragraph(doc, line)
            i += 1

        # Metadata
        doc.core_properties.author = "Sovereign AI Workbench"
        doc.core_properties.title = title

        buf = BytesIO()
        doc.save(buf)
        logger.debug(
            "ExportService: DOCX generated for '%s' (%d bytes)",
            title, buf.tell(),
        )
        return buf.getvalue()

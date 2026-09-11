"""
Tests for the Agent Core.

Covers:
- Tool abstraction and ToolResult
- ToolRegistry
- CalculatorTool (including safety)
- FileReaderTool (including path traversal, unsupported types)
- TaskRouter (deterministic routing)
- TaskPlanner (ExecutionPlan, PlanStep, validation)
- AgentExecutor (sequential execution, tool failure)
- Verifier (calculator verification, NOT_VERIFIED fallback)
- ExecutionTrace (event lifecycle)
- Orchestrator (full pipeline with mocked model)
- POST /agent/run response structure

All tests use mocked HTTP responses and do NOT require a real
llama.cpp server.
"""

import json
from pathlib import Path
from unittest.mock import patch, MagicMock

import pytest
import httpx

from app.tools.base import Tool, ToolResult
from app.tools.registry import ToolRegistry
from app.tools.calculator import CalculatorTool
from app.tools.file_reader import FileReaderTool
from app.agents.router import TaskRouter, RouteResult
from app.agents.planner import TaskPlanner, ExecutionPlan, PlanStep, PlanValidationError
from app.agents.executor import AgentExecutor, ExecutionResult
from app.agents.verifier import (
    Verifier,
    VerifierRegistry,
    CalculatorVerifier,
    VerificationStatus,
    VerificationResult,
)
from app.agents.trace import ExecutionTrace, EventType, TraceEvent
from app.agents.orchestrator import AgentOrchestrator, OrchestratorResult
from app.models.registry import ModelRegistry
from app.models.local import DummyLocalModel
from app.security.audit import AuditService


# ================================================================ helpers


def _mock_chat_response(text: str = "Hello", total_tokens: int = 42):
    return {
        "choices": [
            {"index": 0, "message": {"role": "assistant", "content": text}, "finish_reason": "stop"}
        ],
        "usage": {"prompt_tokens": 10, "completion_tokens": 32, "total_tokens": total_tokens},
    }


# ================================================================ Tool abstraction


class TestToolResult:
    def test_success_result(self) -> None:
        r = ToolResult(success=True, result=42)
        assert r.success is True
        assert r.result == 42
        assert r.error is None

    def test_failure_result(self) -> None:
        r = ToolResult(success=False, error="bad input")
        assert r.success is False
        assert r.error == "bad input"

    def test_metadata(self) -> None:
        r = ToolResult(success=True, result="ok", metadata={"key": "val"})
        assert r.metadata["key"] == "val"


# ================================================================ ToolRegistry


class TestToolRegistry:
    def test_register_and_get(self) -> None:
        reg = ToolRegistry()
        calc = CalculatorTool()
        reg.register(calc)
        assert reg.has("calculator")
        assert reg.get("calculator") is calc

    def test_duplicate_raises(self) -> None:
        reg = ToolRegistry()
        reg.register(CalculatorTool())
        with pytest.raises(ValueError, match="already registered"):
            reg.register(CalculatorTool())

    def test_get_missing_raises(self) -> None:
        reg = ToolRegistry()
        with pytest.raises(KeyError, match="No tool registered"):
            reg.get("nonexistent")

    def test_list_tools(self, tmp_path) -> None:
        reg = ToolRegistry()
        reg.register(CalculatorTool())
        reg.register(FileReaderTool(workspace_root=tmp_path))
        assert set(reg.list_tools()) == {"calculator", "file_reader"}

    def test_has_false(self) -> None:
        reg = ToolRegistry()
        assert reg.has("nope") is False


# ================================================================ CalculatorTool


class TestCalculatorTool:
    def setup_method(self):
        self.calc = CalculatorTool()

    def test_name(self) -> None:
        assert self.calc.name == "calculator"

    def test_description(self) -> None:
        assert "arithmetic" in self.calc.description.lower()

    def test_input_schema(self) -> None:
        schema = self.calc.input_schema
        assert schema["type"] == "object"
        assert "expression" in schema["properties"]

    def test_addition(self) -> None:
        r = self.calc.execute({"expression": "2 + 3"})
        assert r.success is True
        assert r.result == 5

    def test_subtraction(self) -> None:
        r = self.calc.execute({"expression": "10 - 4"})
        assert r.success and r.result == 6

    def test_multiplication(self) -> None:
        r = self.calc.execute({"expression": "125 * 37"})
        assert r.success and r.result == 4625

    def test_division(self) -> None:
        r = self.calc.execute({"expression": "100 / 4"})
        assert r.success and r.result == 25

    def test_floor_division(self) -> None:
        r = self.calc.execute({"expression": "7 // 2"})
        assert r.success and r.result == 3

    def test_modulo(self) -> None:
        r = self.calc.execute({"expression": "10 % 3"})
        assert r.success and r.result == 1

    def test_power(self) -> None:
        r = self.calc.execute({"expression": "2 ** 10"})
        assert r.success and r.result == 1024

    def test_parentheses(self) -> None:
        r = self.calc.execute({"expression": "(2 + 3) * 4"})
        assert r.success and r.result == 20

    def test_negative_number(self) -> None:
        r = self.calc.execute({"expression": "-5 + 3"})
        assert r.success and r.result == -2

    def test_float_result(self) -> None:
        r = self.calc.execute({"expression": "7 / 2"})
        assert r.success and r.result == 3.5

    def test_complex_expression(self) -> None:
        r = self.calc.execute({"expression": "(10 + 20 + 30) / 3"})
        assert r.success and r.result == 20


class TestCalculatorSafety:
    def setup_method(self):
        self.calc = CalculatorTool()

    def test_reject_empty(self) -> None:
        r = self.calc.execute({"expression": ""})
        assert r.success is False

    def test_reject_function_call(self) -> None:
        r = self.calc.execute({"expression": "os.system('rm -rf /')"})
        assert r.success is False

    def test_reject_import(self) -> None:
        r = self.calc.execute({"expression": "__import__('os')"})
        assert r.success is False

    def test_reject_variable_name(self) -> None:
        r = self.calc.execute({"expression": "x + 1"})
        assert r.success is False

    def test_reject_string_literal(self) -> None:
        r = self.calc.execute({"expression": "'hello'"})
        assert r.success is False

    def test_reject_list_comprehension(self) -> None:
        r = self.calc.execute({"expression": "[x for x in range(10)]"})
        assert r.success is False

    def test_reject_huge_exponent(self) -> None:
        r = self.calc.execute({"expression": "2 ** 999999"})
        assert r.success is False
        assert "Exponent too large" in r.error

    def test_division_by_zero(self) -> None:
        r = self.calc.execute({"expression": "1 / 0"})
        assert r.success is False

    def test_reject_semicolon_injection(self) -> None:
        r = self.calc.execute({"expression": "1; import os"})
        assert r.success is False

    def test_reject_lambda(self) -> None:
        r = self.calc.execute({"expression": "lambda: 1"})
        assert r.success is False

    def test_reject_non_string_input(self) -> None:
        r = self.calc.execute({"expression": 42})
        assert r.success is False

    def test_reject_missing_expression(self) -> None:
        r = self.calc.execute({})
        assert r.success is False


# ================================================================ FileReaderTool


class TestFileReaderTool:
    @pytest.fixture()
    def workspace(self, tmp_path) -> Path:
        """Create a workspace with test files."""
        ws = tmp_path / "workspace"
        ws.mkdir()

        (ws / "hello.txt").write_text("Hello World\n", encoding="utf-8")
        (ws / "numbers.txt").write_text("10\n20\n30\n40\n50\n", encoding="utf-8")
        (ws / "data.csv").write_text("a,b,c\n1,2,3\n", encoding="utf-8")

        sub = ws / "subdir"
        sub.mkdir()
        (sub / "nested.txt").write_text("nested content", encoding="utf-8")

        return ws

    @pytest.fixture()
    def reader(self, workspace) -> FileReaderTool:
        return FileReaderTool(workspace_root=workspace)

    def test_name(self, reader) -> None:
        assert reader.name == "file_reader"

    def test_description(self, reader) -> None:
        assert "read" in reader.description.lower()

    def test_read_txt(self, reader) -> None:
        r = reader.execute({"file_path": "hello.txt"})
        assert r.success is True
        assert "Hello World" in r.result

    def test_read_nested(self, reader) -> None:
        r = reader.execute({"file_path": "subdir/nested.txt"})
        assert r.success is True
        assert "nested content" in r.result

    def test_file_not_found(self, reader) -> None:
        r = reader.execute({"file_path": "nonexistent.txt"})
        assert r.success is False
        assert "not found" in r.error.lower()

    def test_unsupported_extension(self, reader) -> None:
        r = reader.execute({"file_path": "data.csv"})
        assert r.success is False
        assert "Unsupported" in r.error

    def test_empty_path(self, reader) -> None:
        r = reader.execute({"file_path": ""})
        assert r.success is False


class TestFileReaderPathTraversal:
    @pytest.fixture()
    def workspace(self, tmp_path) -> Path:
        ws = tmp_path / "workspace"
        ws.mkdir()
        (ws / "safe.txt").write_text("safe", encoding="utf-8")
        # Create a file outside the workspace.
        (tmp_path / "secret.txt").write_text("secret data", encoding="utf-8")
        return ws

    @pytest.fixture()
    def reader(self, workspace) -> FileReaderTool:
        return FileReaderTool(workspace_root=workspace)

    def test_block_parent_traversal(self, reader) -> None:
        r = reader.execute({"file_path": "../secret.txt"})
        assert r.success is False
        assert "Access denied" in r.error

    def test_block_absolute_outside(self, reader, tmp_path) -> None:
        r = reader.execute({"file_path": str(tmp_path / "secret.txt")})
        assert r.success is False
        assert "Access denied" in r.error

    def test_block_double_dot_in_middle(self, reader) -> None:
        r = reader.execute({"file_path": "subdir/../../secret.txt"})
        assert r.success is False
        assert "Access denied" in r.error

    def test_allow_absolute_within_workspace(self, reader, workspace) -> None:
        r = reader.execute({"file_path": str(workspace / "safe.txt")})
        assert r.success is True
        assert r.result == "safe"

    def test_workspace_not_a_dir_raises(self, tmp_path) -> None:
        fake = tmp_path / "not_a_dir"
        fake.write_text("x")
        with pytest.raises(ValueError, match="not a directory"):
            FileReaderTool(workspace_root=fake)


class TestFileReaderDocx:
    def test_read_docx(self, tmp_path) -> None:
        import docx
        doc = docx.Document()
        doc.add_paragraph("Test paragraph one.")
        doc.add_paragraph("Test paragraph two.")
        path = tmp_path / "test.docx"
        doc.save(str(path))

        reader = FileReaderTool(workspace_root=tmp_path)
        r = reader.execute({"file_path": "test.docx"})
        assert r.success is True
        assert "Test paragraph one" in r.result
        assert "Test paragraph two" in r.result


class TestFileReaderPdf:
    def test_read_pdf_with_text(self, tmp_path) -> None:
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        # pypdf blank pages have no text; this tests the "ocr_required" path.
        path = tmp_path / "blank.pdf"
        with open(path, "wb") as f:
            writer.write(f)

        reader = FileReaderTool(workspace_root=tmp_path)
        r = reader.execute({"file_path": "blank.pdf"})
        assert r.success is True
        assert r.metadata.get("status") == "ocr_required"


# ================================================================ TaskRouter


class TestTaskRouter:
    def setup_method(self):
        self.router = TaskRouter()

    def test_calculation_route(self) -> None:
        r = self.router.route("calculate 2 + 2")
        assert r.task_type == "calculation"
        assert "calculator" in r.tools_hint

    def test_file_route(self) -> None:
        r = self.router.route("read the file data.txt")
        assert r.task_type == "file_analysis"
        assert "file_reader" in r.tools_hint

    def test_multi_step_route(self) -> None:
        r = self.router.route("read data.txt and calculate the average")
        assert r.task_type == "multi_step"
        assert "file_reader" in r.tools_hint
        assert "calculator" in r.tools_hint

    def test_general_route(self) -> None:
        r = self.router.route("explain quantum computing")
        assert r.task_type == "general"
        assert r.tools_hint == []

    def test_empty_task(self) -> None:
        r = self.router.route("")
        assert r.task_type == "general"

    def test_compute_keyword(self) -> None:
        r = self.router.route("compute 5 * 10")
        assert r.task_type == "calculation"

    def test_what_is_pattern(self) -> None:
        r = self.router.route("what is 100 / 4")
        assert r.task_type == "calculation"

    def test_file_with_connector(self) -> None:
        r = self.router.route("read report.pdf and then summarize")
        assert r.task_type == "multi_step"

    def test_returns_route_result(self) -> None:
        r = self.router.route("calculate 1+1")
        assert isinstance(r, RouteResult)
        assert r.confidence > 0


# ================================================================ TaskPlanner


class TestTaskPlanner:
    @pytest.fixture()
    def tool_reg(self, tmp_path) -> ToolRegistry:
        reg = ToolRegistry()
        reg.register(CalculatorTool())
        reg.register(FileReaderTool(workspace_root=tmp_path))
        return reg

    @pytest.fixture()
    def planner(self, tool_reg) -> TaskPlanner:
        return TaskPlanner(tool_reg)

    def test_plan_calculation(self, planner) -> None:
        route = RouteResult(task_type="calculation", tools_hint=["calculator"])
        plan = planner.plan("calculate 2 + 2", route)
        assert isinstance(plan, ExecutionPlan)
        assert len(plan.steps) == 1
        assert plan.steps[0].tool == "calculator"
        assert "2 + 2" in plan.steps[0].input.get("expression", "")

    def test_plan_file_analysis(self, planner) -> None:
        route = RouteResult(task_type="file_analysis", tools_hint=["file_reader"])
        plan = planner.plan("read data.txt", route)
        assert len(plan.steps) == 1
        assert plan.steps[0].tool == "file_reader"
        assert "data.txt" in plan.steps[0].input.get("file_path", "")

    def test_plan_multi_step(self, planner) -> None:
        route = RouteResult(
            task_type="multi_step",
            tools_hint=["file_reader", "calculator"],
        )
        plan = planner.plan("read data.txt and calculate the average", route)
        assert len(plan.steps) == 2
        assert plan.steps[0].tool == "file_reader"
        assert plan.steps[1].tool == "calculator"
        assert len(plan.steps[1].depends_on) > 0

    def test_plan_general_no_steps(self, planner) -> None:
        route = RouteResult(task_type="general", tools_hint=[])
        plan = planner.plan("explain quantum computing", route)
        assert len(plan.steps) == 0

    def test_plan_has_run_id(self, planner) -> None:
        route = RouteResult(task_type="calculation", tools_hint=["calculator"])
        plan = planner.plan("calculate 1+1", route)
        assert plan.run_id is not None
        assert len(plan.run_id) > 0

    def test_step_has_step_id(self, planner) -> None:
        route = RouteResult(task_type="calculation", tools_hint=["calculator"])
        plan = planner.plan("calculate 1+1", route)
        assert plan.steps[0].step_id is not None


class TestPlanValidation:
    def test_reject_unregistered_tool(self) -> None:
        reg = ToolRegistry()  # Empty — no tools.
        planner = TaskPlanner(reg)
        route = RouteResult(task_type="calculation", tools_hint=["calculator"])
        with pytest.raises(PlanValidationError, match="not registered"):
            planner.plan("calculate 1+1", route)


# ================================================================ AgentExecutor


class TestAgentExecutor:
    @pytest.fixture()
    def tool_reg(self, tmp_path) -> ToolRegistry:
        reg = ToolRegistry()
        reg.register(CalculatorTool())
        ws = tmp_path / "workspace"
        ws.mkdir()
        (ws / "numbers.txt").write_text("10\n20\n30\n", encoding="utf-8")
        reg.register(FileReaderTool(workspace_root=ws))
        return reg

    @pytest.fixture()
    def verifier_reg(self) -> VerifierRegistry:
        vr = VerifierRegistry()
        vr.register(CalculatorVerifier())
        return vr

    @pytest.fixture()
    def executor(self, tool_reg, verifier_reg) -> AgentExecutor:
        return AgentExecutor(tool_reg, verifier_reg)

    def test_execute_single_calculator_step(self, executor) -> None:
        plan = ExecutionPlan(
            task="calculate 2+3",
            steps=[PlanStep(tool="calculator", description="calc", input={"expression": "2+3"})],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is True
        assert len(result.outcomes) == 1
        assert result.outcomes[0].tool_result.result == 5

    def test_execute_single_file_step(self, executor) -> None:
        plan = ExecutionPlan(
            task="read numbers.txt",
            steps=[PlanStep(tool="file_reader", description="read", input={"file_path": "numbers.txt"})],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is True
        assert "10" in result.outcomes[0].tool_result.result

    def test_execute_stops_on_failure(self, executor) -> None:
        plan = ExecutionPlan(
            task="fail",
            steps=[
                PlanStep(tool="file_reader", description="read missing", input={"file_path": "missing.txt"}),
                PlanStep(tool="calculator", description="never reached", input={"expression": "1+1"}),
            ],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is False
        assert len(result.outcomes) == 1  # Second step never executed.

    def test_execute_unknown_tool(self, executor) -> None:
        plan = ExecutionPlan(
            task="test",
            steps=[PlanStep(tool="nonexistent", description="bad", input={})],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is False
        assert "not registered" in result.outcomes[0].tool_result.error

    def test_verification_runs_after_calculator(self, executor) -> None:
        plan = ExecutionPlan(
            task="calculate",
            steps=[PlanStep(tool="calculator", description="calc", input={"expression": "3*7"})],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is True
        v = result.outcomes[0].verification
        assert v is not None
        assert v.status == VerificationStatus.PASS

    def test_verification_not_verified_for_file_reader(self, executor) -> None:
        plan = ExecutionPlan(
            task="read",
            steps=[PlanStep(tool="file_reader", description="read", input={"file_path": "numbers.txt"})],
        )
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)
        assert result.success is True
        v = result.outcomes[0].verification
        assert v is not None
        assert v.status == VerificationStatus.NOT_VERIFIED


class TestAgentExecutorMultiStep:
    """Test the multi-step file_reader -> calculator pipeline."""

    @pytest.fixture()
    def tool_reg(self, tmp_path) -> ToolRegistry:
        reg = ToolRegistry()
        reg.register(CalculatorTool())
        ws = tmp_path / "workspace"
        ws.mkdir()
        (ws / "data.txt").write_text("10\n20\n30\n", encoding="utf-8")
        reg.register(FileReaderTool(workspace_root=ws))
        return reg

    @pytest.fixture()
    def executor(self, tool_reg) -> AgentExecutor:
        vr = VerifierRegistry()
        vr.register(CalculatorVerifier())
        return AgentExecutor(tool_reg, vr)

    def test_multi_step_file_then_calculator(self, executor) -> None:
        step1 = PlanStep(
            step_id="s1",
            tool="file_reader",
            description="Read file",
            input={"file_path": "data.txt"},
        )
        step2 = PlanStep(
            step_id="s2",
            tool="calculator",
            description="Calculate average",
            input={"expression": ""},  # Will be populated from file content.
            depends_on=["s1"],
        )
        plan = ExecutionPlan(task="read and average", steps=[step1, step2])
        trace = ExecutionTrace(run_id=plan.run_id)
        result = executor.execute(plan, trace)

        assert result.success is True
        assert len(result.outcomes) == 2
        # File was read.
        assert result.outcomes[0].tool_result.success is True
        # Calculator computed average of 10, 20, 30 = 20.
        assert result.outcomes[1].tool_result.success is True
        assert result.outcomes[1].tool_result.result == 20


# ================================================================ Verifier


class TestVerifier:
    def test_calculator_verifier_pass(self) -> None:
        v = CalculatorVerifier()
        tool_input = {"expression": "6 * 7"}
        tool_result = ToolResult(success=True, result=42)
        vr = v.verify(tool_input, tool_result)
        assert vr.status == VerificationStatus.PASS

    def test_calculator_verifier_fail(self) -> None:
        v = CalculatorVerifier()
        tool_input = {"expression": "6 * 7"}
        tool_result = ToolResult(success=True, result=999)  # Wrong!
        vr = v.verify(tool_input, tool_result)
        assert vr.status == VerificationStatus.FAIL

    def test_calculator_verifier_not_verified_on_error(self) -> None:
        v = CalculatorVerifier()
        tool_input = {"expression": "6 * 7"}
        tool_result = ToolResult(success=False, error="some error")
        vr = v.verify(tool_input, tool_result)
        assert vr.status == VerificationStatus.NOT_VERIFIED

    def test_verifier_registry_returns_not_verified_for_unknown(self) -> None:
        vr = VerifierRegistry()
        result = vr.verify("unknown_tool", {}, ToolResult(success=True, result="x"))
        assert result.status == VerificationStatus.NOT_VERIFIED

    def test_verifier_registry_dispatches(self) -> None:
        vr = VerifierRegistry()
        vr.register(CalculatorVerifier())
        result = vr.verify("calculator", {"expression": "1+1"}, ToolResult(success=True, result=2))
        assert result.status == VerificationStatus.PASS


# ================================================================ ExecutionTrace


class TestExecutionTrace:
    def test_emit_event(self) -> None:
        trace = ExecutionTrace(run_id="test-run")
        event = trace.emit(EventType.TASK_RECEIVED, metadata={"task": "test"})
        assert event.run_id == "test-run"
        assert event.event_type == EventType.TASK_RECEIVED
        assert len(trace.events) == 1

    def test_event_has_timestamp(self) -> None:
        trace = ExecutionTrace()
        event = trace.emit(EventType.TASK_RECEIVED)
        assert event.timestamp is not None

    def test_lifecycle_events(self) -> None:
        trace = ExecutionTrace()
        trace.emit(EventType.TASK_RECEIVED)
        trace.emit(EventType.TASK_ROUTED)
        trace.emit(EventType.PLAN_CREATED)
        trace.emit(EventType.TOOL_STARTED, step_id="s1")
        trace.emit(EventType.TOOL_COMPLETED, step_id="s1")
        trace.emit(EventType.VERIFICATION_STARTED, step_id="s1")
        trace.emit(EventType.VERIFICATION_COMPLETED, step_id="s1")
        trace.emit(EventType.TASK_COMPLETED)
        assert len(trace.events) == 8
        types = [e.event_type for e in trace.events]
        assert types[0] == EventType.TASK_RECEIVED
        assert types[-1] == EventType.TASK_COMPLETED

    def test_to_dicts(self) -> None:
        trace = ExecutionTrace(run_id="r1")
        trace.emit(EventType.TASK_RECEIVED)
        dicts = trace.to_dicts()
        assert len(dicts) == 1
        assert dicts[0]["run_id"] == "r1"
        assert dicts[0]["event_type"] == "TASK_RECEIVED"

    def test_step_id_attached(self) -> None:
        trace = ExecutionTrace()
        event = trace.emit(EventType.TOOL_STARTED, step_id="step-42")
        assert event.step_id == "step-42"

    def test_auto_run_id(self) -> None:
        trace = ExecutionTrace()
        assert trace.run_id is not None
        assert len(trace.run_id) > 0


# ================================================================ Orchestrator (full pipeline)


class TestOrchestratorPipeline:
    """Test the full orchestrator pipeline with DummyLocalModel."""

    @pytest.fixture()
    def orch(self, tmp_path) -> AgentOrchestrator:
        model_reg = ModelRegistry()
        model_reg.register("general", DummyLocalModel())
        audit = AuditService(log_file=tmp_path / "audit.log")

        tool_reg = ToolRegistry()
        tool_reg.register(CalculatorTool())

        ws = tmp_path / "workspace"
        ws.mkdir()
        (ws / "numbers.txt").write_text("10\n20\n30\n", encoding="utf-8")
        tool_reg.register(FileReaderTool(workspace_root=ws))

        verifier_reg = VerifierRegistry()
        verifier_reg.register(CalculatorVerifier())

        return AgentOrchestrator(
            registry=model_reg,
            audit_service=audit,
            tool_registry=tool_reg,
            verifier_registry=verifier_reg,
        )

    def test_general_task(self, orch) -> None:
        result = orch.run("explain quantum computing")
        assert result.execution_status == "success"
        assert result.task_type == "general"
        assert len(result.plan) == 0
        assert len(result.tool_calls) == 0

    def test_calculation_task(self, orch) -> None:
        result = orch.run("calculate 125 * 37")
        assert result.execution_status == "success"
        assert result.task_type == "calculation"
        assert len(result.plan) == 1
        assert result.plan[0]["tool"] == "calculator"
        assert len(result.tool_calls) == 1
        assert result.tool_calls[0]["success"] is True
        assert result.tool_calls[0]["verification"] == "PASS"

    def test_file_reading_task(self, orch) -> None:
        result = orch.run("read numbers.txt")
        assert result.execution_status == "success"
        assert result.task_type == "file_analysis"
        assert len(result.plan) == 1
        assert result.plan[0]["tool"] == "file_reader"

    def test_multi_step_task(self, orch) -> None:
        result = orch.run("read numbers.txt and calculate the average")
        assert result.execution_status == "success"
        assert result.task_type == "multi_step"
        assert len(result.plan) == 2
        assert result.plan[0]["tool"] == "file_reader"
        assert result.plan[1]["tool"] == "calculator"

    def test_result_has_run_id(self, orch) -> None:
        result = orch.run("calculate 1+1")
        assert result.run_id is not None
        assert len(result.run_id) > 0

    def test_result_has_verification_summary(self, orch) -> None:
        result = orch.run("calculate 5 * 5")
        assert "verified_steps" in result.verification

    def test_result_has_trace(self, orch) -> None:
        result = orch.run("calculate 2 + 2")
        assert len(result.trace) > 0
        types = [e["event_type"] for e in result.trace]
        assert "TASK_RECEIVED" in types
        assert "TASK_COMPLETED" in types

    def test_missing_model_error(self, tmp_path) -> None:
        empty_reg = ModelRegistry()
        audit = AuditService(log_file=tmp_path / "audit.log")
        orch = AgentOrchestrator(registry=empty_reg, audit_service=audit)
        result = orch.run("test")
        assert result.execution_status == "error"
        assert "Model selection failed" in result.result


class TestOrchestratorWithLlamaCppMocked:
    """Ensure the orchestrator still works with LlamaCppProvider (mocked)."""

    def test_calculation_with_llama_cpp(self, tmp_path) -> None:
        from app.models.llama_cpp_provider import LlamaCppProvider

        model_reg = ModelRegistry()
        provider = LlamaCppProvider(base_url="http://127.0.0.1:8080")
        model_reg.register("general", provider)
        audit = AuditService(log_file=tmp_path / "audit.log")

        tool_reg = ToolRegistry()
        tool_reg.register(CalculatorTool())

        verifier_reg = VerifierRegistry()
        verifier_reg.register(CalculatorVerifier())

        orch = AgentOrchestrator(
            registry=model_reg,
            audit_service=audit,
            tool_registry=tool_reg,
            verifier_registry=verifier_reg,
        )

        mock_resp = MagicMock()
        mock_resp.status_code = 200
        mock_resp.json.return_value = _mock_chat_response("The result is 4625.", 50)
        mock_health = MagicMock()
        mock_health.status_code = 200

        with patch.object(httpx, "get", return_value=mock_health), \
             patch.object(httpx, "post", return_value=mock_resp):
            result = orch.run("calculate 125 * 37")

        assert result.execution_status == "success"
        assert result.provider == "llama_cpp"
        assert result.task_type == "calculation"
        assert len(result.tool_calls) == 1
        assert result.tool_calls[0]["success"] is True


# ================================================================ API response structure


class TestAgentRunAPIStructure:
    @pytest.fixture()
    def client(self):
        from fastapi.testclient import TestClient
        from app.main import create_app
        app = create_app()
        with TestClient(app) as c:
            yield c

    def test_response_has_new_fields(self, client) -> None:
        resp = client.post("/agent/run", json={"task": "explain AI"})
        assert resp.status_code == 200
        data = resp.json()
        assert "run_id" in data
        assert "task_type" in data
        assert "plan" in data
        assert "tool_calls" in data
        assert "verification" in data
        assert "result" in data
        assert "provider" in data
        assert "execution_status" in data

    def test_calculation_response_structure(self, client) -> None:
        resp = client.post("/agent/run", json={"task": "calculate 10 + 20"})
        data = resp.json()
        assert data["task_type"] == "calculation"
        assert data["execution_status"] == "success"
        assert len(data["plan"]) == 1
        assert data["plan"][0]["tool"] == "calculator"
        assert len(data["tool_calls"]) == 1
        assert data["tool_calls"][0]["success"] is True

    def test_empty_task_still_rejected(self, client) -> None:
        resp = client.post("/agent/run", json={"task": "   "})
        assert resp.status_code == 422

    def test_general_task_response(self, client) -> None:
        resp = client.post("/agent/run", json={"task": "what is the meaning of life"})
        data = resp.json()
        assert data["task_type"] == "general"
        assert data["execution_status"] == "success"
        assert len(data["plan"]) == 0
        assert len(data["tool_calls"]) == 0

    def test_health_includes_tools(self, client) -> None:
        data = client.get("/health").json()
        assert "tools_registered" in data
        assert "calculator" in data["tools_registered"]
        assert "file_reader" in data["tools_registered"]


# ================================================================ Audit integration


class TestAuditEventLifecycle:
    def test_audit_records_task_type(self, tmp_path) -> None:
        model_reg = ModelRegistry()
        model_reg.register("general", DummyLocalModel())
        audit = AuditService(log_file=tmp_path / "audit.log")

        tool_reg = ToolRegistry()
        tool_reg.register(CalculatorTool())

        orch = AgentOrchestrator(
            registry=model_reg,
            audit_service=audit,
            tool_registry=tool_reg,
        )
        orch.run("calculate 1+1")

        records = audit.get_recent()
        assert len(records) >= 1
        last = records[-1]
        assert last.metadata.get("task_type") == "calculation"
        assert last.metadata.get("local_inference") is True

    def test_audit_records_tools_used(self, tmp_path) -> None:
        model_reg = ModelRegistry()
        model_reg.register("general", DummyLocalModel())
        audit = AuditService(log_file=tmp_path / "audit.log")

        tool_reg = ToolRegistry()
        tool_reg.register(CalculatorTool())

        orch = AgentOrchestrator(
            registry=model_reg,
            audit_service=audit,
            tool_registry=tool_reg,
        )
        orch.run("calculate 5*5")

        records = audit.get_recent()
        last = records[-1]
        assert "calculator" in last.metadata.get("tools_used", [])
